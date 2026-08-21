"""
사무소별 매핑 파일 관리 모듈
mappings/*.json → 열 매핑 로드 → 엑셀 입력
"""

import json
from pathlib import Path
from copy import copy
import openpyxl

from core.schema import to_excel_value

MAPPINGS_DIR = Path(__file__).parent.parent / "mappings"


# ─── 매핑 파일 목록 ───────────────────────────────────────────────────────────

# ── 필드명 별칭 ──────────────────────────────────────────────────────────────
#  같은 항목을 사무소마다 다르게 부른다. 실제 매핑 파일 기준:
#    초본/인감            → 검단롯데캐슬넥스티엘, 남양주 자연앤이편한세상3차
#    초본발급일/인감발급일 → 검단롯데캐슬넥스티엘_기초입력
#    초본발행일/인감발행일 → 법무법인제이엘, 한예라_남양주
#  추출기는 이 중 일부만 만들기 때문에, 예전에는 표기가 다른 사무소의
#  초본·인감 열이 **영구히 빈칸**이었다. 기입 직전에 별칭을 해석해
#  어느 표기를 쓰든 값이 채워지도록 한다.
FIELD_ALIASES = {
    "초본발행일": ["초본발급일", "초본"],
    "초본발급일": ["초본발행일", "초본"],
    "초본":       ["초본발행일", "초본발급일"],
    "인감발행일": ["인감발급일", "인감"],
    "인감발급일": ["인감발행일", "인감"],
    "인감":       ["인감발행일", "인감발급일"],
    "등본발행일": ["등본발급일", "등본"],
    "등본발급일": ["등본발행일", "등본"],
    "등본":       ["등본발행일", "등본발급일"],
    "거래신고필증번호": ["실거래일련번호"],
    "실거래일련번호":   ["거래신고필증번호"],
    "전용면적": ["건물면적"],
    "대지지분": ["대지권면적"],
}


def _resolve_value(record: dict, field: str):
    """매핑이 요구하는 필드명으로 값을 찾되, 없으면 알려진 별칭으로 재시도."""
    v = record.get(field)
    if v not in (None, ""):
        return v
    for alt in FIELD_ALIASES.get(field, ()):
        v = record.get(alt)
        if v not in (None, ""):
            return v
    return None


def _is_formula_cell(cell) -> bool:
    """
    셀이 수식인지 판정.
    일반 수식은 '=' 로 시작하는 str 이지만, 배열수식/데이터테이블은
    openpyxl 이 ArrayFormula / DataTableFormula 객체로 반환하므로
    str 검사만으로는 놓쳐 수식을 덮어쓰게 된다. 두 경우 모두 잡는다.
    """
    v = cell.value
    if isinstance(v, str):
        return v.startswith("=")
    if v is None:
        return False
    # ArrayFormula / DataTableFormula 등 openpyxl 수식 객체
    return type(v).__name__ in ("ArrayFormula", "DataTableFormula")


def _is_writable(cell) -> bool:
    """
    병합셀의 비-좌상단 셀(MergedCell)은 value 대입 시 AttributeError 를 던진다.
    write_with_mapping 에는 try/except 가 없어 그 한 번으로 저장 전체가 실패하고
    한 건도 저장되지 않는다. 미리 걸러낸다.
    """
    return type(cell).__name__ != "MergedCell"


def _write_cell(ws, row, col, val, is_amount=False, field: str = None):
    """
    셀에 값 쓰기 — 필드 타입(schema.to_excel_value)에 맞춰 실제 파이썬
    타입(int/float/date)으로 변환 후 기입한다.
    is_amount=True 이면(매핑 _amount_columns 지정) 필드 타입과 무관하게
    숫자로 강제 변환한다.
    """
    from core.schema import to_excel_value
    force = "int" if is_amount else None
    ws.cell(row=row, column=col).value = to_excel_value(field, val, force_type=force)


def _clear_data_rows(ws, data_start: int, formula_cols: set = None,
                     preserve_cols: set = None, target_cols: set = None,
                     key_col: int = None, max_scan: int = 20000):
    """
    새로쓰기 전 기존 데이터를 비운다.

    이전 구현의 3가지 치명적 결함을 수정:
    1) 매핑에 없는 열까지 전부 삭제했다. 실제 템플릿은 열이 200개가 넘어
       (template.xlsx 는 223열) 사무장이 손으로 적어둔 메모·특이사항·
       담당자 등 **매핑 밖 모든 수기 입력이 소리 없이 증발**했다.
       → target_cols(매핑이 실제로 쓰는 열)로 삭제 범위를 한정한다.
    2) 중간에 빈 행(또는 수식만 있는 행)을 만나면 break 해서, 그 아래
       지난번 명단 잔해가 그대로 남았다. 수식이 미리 깔린 템플릿에서는
       거의 항상 발생한다.
       → break 하지 않고 연속 빈 행이 충분히 이어질 때만 종료한다.
    3) data_start+1000 상한 때문에 1000세대 초과 단지는 아래가 안 지워졌다.
       → max_scan 으로 넉넉히 확장(기본 20000행).

    preserve_cols 도 이제 실제로 전달받아 보존한다.
    """
    formula_cols  = formula_cols or set()
    preserve_cols = preserve_cols or set()
    보호 = formula_cols | preserve_cols

    # 삭제 대상 열: 매핑이 쓰는 열만. 지정 없으면 (구 동작 호환) 전체.
    if target_cols:
        대상열 = [c for c in sorted(target_cols) if c not in 보호]
    else:
        대상열 = [c for c in range(1, ws.max_column + 1) if c not in 보호]

    검사열 = [key_col] if key_col else list(range(1, min(ws.max_column, 40) + 1))

    끝행 = min(ws.max_row, data_start + max_scan)
    연속빈행 = 0
    for r in range(data_start, 끝행 + 1):
        # 이 행에 '실데이터'(수식 아닌 값)가 있는지 — 검사열만 훑어 성능 확보
        row_has_data = False
        for c in 검사열:
            cell = ws.cell(r, c)
            v = cell.value
            if v is not None and v != "" and not _is_formula_cell(cell):
                row_has_data = True
                break

        if not row_has_data:
            연속빈행 += 1
            # 충분히 이어지면 그 아래는 비었다고 보고 종료 (중간 빈 행 1~2줄은 통과)
            if 연속빈행 >= 50:
                break
            continue
        연속빈행 = 0

        for c in 대상열:
            cell = ws.cell(r, c)
            if _is_formula_cell(cell) or not _is_writable(cell):
                continue   # 수식(배열수식 포함) / 병합셀 보존
            cell.value = None


def get_mapping_list() -> list[str]:
    """사무소 이름 목록 반환 (_presets 제외)"""
    files = sorted(MAPPINGS_DIR.glob("*.json"))
    return [f.stem for f in files if not f.stem.startswith("_")]


def load_mapping(사무소명: str) -> dict:
    """사무소명으로 매핑 로드"""
    path = MAPPINGS_DIR / f"{사무소명}.json"
    if not path.exists():
        raise FileNotFoundError(f"매핑 파일 없음: {path}")
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save_mapping(사무소명: str, mapping: dict):
    """매핑 저장"""
    MAPPINGS_DIR.mkdir(exist_ok=True)
    path = MAPPINGS_DIR / f"{사무소명}.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(mapping, f, ensure_ascii=False, indent=2)
    print(f"✅ 매핑 저장: {path}")


def delete_mapping(사무소명: str):
    """매핑 삭제"""
    path = MAPPINGS_DIR / f"{사무소명}.json"
    if path.exists():
        path.unlink()


# ─── 엑셀 입력 (매핑 기반) ───────────────────────────────────────────────────

def _atomic_save(wb, output_path: str):
    """임시 파일에 저장 후 원자적 교체 — 저장 중 크래시로 원본이 깨지는 것 방지."""
    import os, tempfile
    d = os.path.dirname(os.path.abspath(output_path)) or "."
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=d)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, output_path)
    except Exception:
        if os.path.exists(tmp):
            try:
                os.remove(tmp)
            except OSError:
                pass
        raise


def _make_backup(path: str):
    """기존 파일을 타임스탬프 백업으로 복사. 없으면 None."""
    import os, shutil
    from datetime import datetime
    if not os.path.exists(path):
        return None
    p = Path(path)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    bak = p.with_name(f"{p.stem}_백업_{stamp}{p.suffix}")
    shutil.copy2(path, bak)
    return str(bak)


def _unit_key(dong, ho) -> tuple:
    """
    동/호 매칭 키 정규화.
    엑셀에 숫자 603 으로 저장돼 있고 추출값이 "0603" 또는 603.0 이면
    문자열 비교가 어긋나 같은 세대가 새 행으로 중복 추가된다.
    앞 0 제거 + float 정수화로 표기 차이를 흡수한다.
    """
    def norm(v):
        s = str(v).strip()
        if s.endswith(".0"):        # 603.0 → 603
            s = s[:-2]
        s2 = s.lstrip("0")          # 0603 → 603
        return s2 if s2 else s
    return (norm(dong), norm(ho))


def _index_existing_units(ws, dong_col: int, ho_col: int, data_start: int) -> dict:
    """기존 행을 스캔해 (동,호) → 행번호 인덱스 구성 (이어쓰기 upsert용)."""
    idx = {}
    # ws.max_row 가 104만으로 잡히는 템플릿이 흔하므로 상한/조기종료 필수
    max_row = min(ws.max_row, data_start + _SCAN_LIMIT)
    빈행 = 0
    for r in range(data_start, max_row + 1):
        dv = ws.cell(row=r, column=dong_col).value
        hv = ws.cell(row=r, column=ho_col).value
        if dv not in (None, "") and hv not in (None, ""):
            idx[_unit_key(dv, hv)] = r
            빈행 = 0
        else:
            빈행 += 1
            if 빈행 >= _EMPTY_RUN_STOP:
                break
    return idx


def write_with_mapping(template_path: str, records: list[dict],
                       사무소명: str, output_path: str = None,
                       append: bool = False, backup: bool = False):
    """
    사무소별 매핑으로 기본명단 엑셀에 데이터 입력.

    append=False (기본): 기존 데이터 클리어 후 처음부터 기입 (단일 실행 저장).
    append=True:         기존 데이터 보존하고 이어서 기입 —
                         순차 모드는 다음 빈 행부터 연번 이어서,
                         동/호 매칭 모드는 같은 동/호는 갱신·새 동/호는 추가(upsert).
    backup=True:         저장 전 기존 output 파일을 타임스탬프 백업.
    저장은 항상 원자적(임시파일 → 교체).
    """
    mapping   = load_mapping(사무소명)
    info      = mapping["_info"]
    col_map   = {k: int(v) for k, v in mapping["_columns"].items()}
    amount_cols = set(mapping.get("_amount_columns", []))
    key_col   = int(mapping.get("_key_column", 1))

    output_path = output_path or template_path

    # 이어쓰기는 기존 output 파일을 기준으로 열어야 데이터가 보존됨
    load_path = output_path if (append and Path(output_path).exists()) else template_path
    if backup:
        _make_backup(output_path)

    wb = openpyxl.load_workbook(load_path)

    # ── 시트 찾기 ────────────────────────────────────────────────────
    # 예전에는 못 찾으면 wb.active 로 폴백했다. 그 결과 사용자가 다른
    # 통합문서를 고르면 마지막에 열려 있던 아무 시트(예: "우리", "수임표")에
    # 32개 열 데이터를 그대로 써서 원본을 되돌릴 수 없게 훼손했다.
    # → 폴백을 없애고, 못 찾으면 명확한 오류로 중단한다.
    sheet_name = info.get("시트명", "")
    ws = None
    if sheet_name and sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    if ws is None:
        후보 = [s for s in wb.sheetnames if "기본명단" in s]
        if len(후보) == 1:
            ws = wb[후보[0]]
        elif len(후보) > 1:
            raise ValueError(
                f"[{사무소명}] 어느 시트에 기입할지 확정할 수 없습니다.\n"
                f"'기본명단'이 들어간 시트가 여러 개입니다: {후보}\n"
                f"매핑의 _info.시트명 을 정확히 지정하세요.")
    if ws is None:
        raise ValueError(
            f"[{사무소명}] 기입할 시트를 찾지 못했습니다.\n"
            f"매핑이 기대한 시트명: '{sheet_name or '(미지정)'}'\n"
            f"이 파일의 시트 목록: {wb.sheetnames}\n"
            f"→ 올바른 기본명단 파일인지 확인하거나, 매핑의 _info.시트명 을 "
            f"이 파일의 실제 시트명으로 수정하세요.")

    print(f"  📋 입력 시트: [{ws.title}]  (모드: {'이어쓰기' if append else '새로쓰기'})")

    data_start = int(info.get("데이터시작행", 2))
    serial_col  = col_map.get("연번")

    # ── 동/호 매칭 모드 (검단 기초입력 등) ────────────────────────
    key_match = mapping.get("_key_match", {})  # {"동": 11, "호": 12}
    formula_cols = set(mapping.get("_formula_cols", []))
    preserve_cols = set(mapping.get("_preserve_cols", []))

    if key_match:
        dong_col = int(key_match.get("동", 0))
        ho_col   = int(key_match.get("호", 0))

        if append:
            # 기존 데이터 보존 + 기존 동/호 인덱스 구성 (upsert)
            row_index = _index_existing_units(ws, dong_col, ho_col, data_start)
        else:
            # 새로쓰기: 기존 데이터 클리어 (매핑이 쓰는 열만, 수식·보존열 제외)
            _clear_data_rows(ws, data_start, formula_cols,
                             preserve_cols=preserve_cols,
                             target_cols=set(col_map.values()) | {dong_col, ho_col},
                             key_col=dong_col or None)
            row_index = {}

        for record in records:
            dong = str(record.get("동", "")).strip()
            ho   = str(record.get("호", "")).strip()
            if not dong or not ho:
                continue
            키 = _unit_key(dong, ho)
            row = row_index.get(키)
            if not row:
                row = _find_next_row(ws, dong_col, data_start)
                ws.cell(row=row, column=dong_col).value = dong
                ws.cell(row=row, column=ho_col).value = ho
                row_index[키] = row

            for field, col_idx in col_map.items():
                if col_idx in formula_cols or col_idx in preserve_cols:
                    continue
                val = _resolve_value(record, field)
                if val is None or val == "":
                    continue
                cell = ws.cell(row=row, column=col_idx)
                if _is_formula_cell(cell) or not _is_writable(cell):
                    continue
                _write_cell(ws, row, col_idx, val, col_idx in amount_cols, field=field)

        _atomic_save(wb, output_path)
        print(f"✅ [{사무소명}] 저장 완료: {output_path}  ({len(records)}건)")
        return output_path

    # ── 순차 모드 ────────────────────────────────────────────────
    if append:
        # 이어쓰기: 기존 데이터 유지, 다음 빈 행부터, 연번 이어서
        start_row   = _find_next_row(ws, serial_col or key_col, data_start)
        last_serial = _get_last_serial(ws, serial_col, data_start) if serial_col else 0
    else:
        # 새로쓰기: 기존 데이터 클리어 (매핑이 쓰는 열만, 수식·보존열 제외)
        _clear_data_rows(ws, data_start, formula_cols,
                         preserve_cols=preserve_cols,
                         target_cols=set(col_map.values()),
                         key_col=(key_col if key_col else None))
        start_row   = data_start
        last_serial = 0

    for i, record in enumerate(records):
        row = start_row + i

        # 위 행 서식 복사 (데이터 있는 마지막 행 기준)
        ref_row = start_row - 1 if start_row > data_start else data_start
        _copy_style(ws, ref_row, row)

        # 연번 — 기존 마지막 연번 이어서 채번.
        # last_serial == -1 이면 연번 열이 수식(=ROW()-1 등)이므로 손대지 않는다.
        if (serial_col and last_serial >= 0
                and serial_col not in formula_cols
                and serial_col not in preserve_cols
                and not _is_formula_cell(ws.cell(row=row, column=serial_col))):
            ws.cell(row=row, column=serial_col).value = last_serial + i + 1

        # 데이터 입력 — 필드 타입에 맞춰 실제 파이썬 타입(int/float/date)으로
        # 변환해서 기입한다. (예전엔 금액열이 아니면 무조건 str() 강제 →
        # 날짜·면적 등이 텍스트로 들어가 다운스트림 수식이 깨지던 버그)
        #
        # ★ 수식 보호: 예전 순차 모드에는 이 보호가 아예 없어서, 템플릿에
        #   미리 깔린 조회 수식(예: 전용면적/대지지분/거래가액을 '주택' 시트에서
        #   INDEX/MATCH 로 끌어오는 열)을 값으로 덮어써 영구 파괴했다.
        #   key_match 모드에만 있던 보호를 순차 모드에도 동일 적용한다.
        for field, col_idx in col_map.items():
            if field in ("연번",):
                continue
            if col_idx in formula_cols or col_idx in preserve_cols:
                continue
            value = _resolve_value(record, field)
            if value is None or value == "":
                continue
            cell = ws.cell(row=row, column=col_idx)
            if _is_formula_cell(cell) or not _is_writable(cell):
                continue   # 템플릿 수식 / 병합셀 보존
            force = "int" if col_idx in amount_cols else None
            cell.value = to_excel_value(field, value, force_type=force)

    # 저장 전 기본명단 시트를 활성(첫 화면)으로 설정
    try:
        wb.active = wb.index(ws)
    except Exception:
        try:
            wb.active = list(wb.sheetnames).index(ws.title)
        except:
            pass

    _atomic_save(wb, output_path)
    print(f"✅ [{사무소명}] 저장 완료: {output_path}  ({len(records)}건)")
    return output_path


#  실제 템플릿에는 ws.max_row 가 1,048,576 으로 잡히는 잔재 행이 흔하다
#  (template.xlsx 실측: max_row=1048576). 전 범위를 훑으면 셀 객체가 수백만 개
#  생성되어 저장이 사실상 멈춘다. 아래 상한/조기종료로 방어한다.
_SCAN_LIMIT = 20000       # 최대 스캔 행 (2만 세대까지 충분)
_EMPTY_RUN_STOP = 50      # 연속 빈 행이 이만큼이면 그 아래는 비었다고 판단


def _find_next_row(ws, key_col: int, start: int) -> int:
    """
    마지막 데이터 행 다음 행 반환.
    수식 셀은 '데이터'로 보지 않는다 — 템플릿에 수식만 미리 깔린 행을
    데이터로 오인하면 명단이 엉뚱하게 아래에서 시작된다.
    """
    last = start - 1
    max_row = min(ws.max_row, start + _SCAN_LIMIT)
    빈행 = 0
    for row in range(start, max_row + 1):
        cell = ws.cell(row=row, column=key_col)
        v = cell.value
        if v not in (None, "") and not _is_formula_cell(cell):
            last = row
            빈행 = 0
        else:
            빈행 += 1
            if 빈행 >= _EMPTY_RUN_STOP:
                break
    return last + 1


def _get_last_serial(ws, serial_col: int, start: int) -> int:
    """
    기존 파일의 마지막 연번 반환 (없으면 0).

    주의: 연번 열이 수식(예: =ROW()-1)이면 셀 값은 문자열이라 int() 가 실패해
    0 이 되고, 이어쓰기 시 1,2,3... 으로 **기존 연번과 중복**된다.
    이 경우 연번은 수식이 스스로 계산하므로 채번 자체를 건너뛰도록
    -1 (=채번 불필요) 을 돌려준다.
    """
    last = 0
    max_row = min(ws.max_row, start + _SCAN_LIMIT)
    빈행 = 0
    수식발견 = False
    for row in range(start, max_row + 1):
        cell = ws.cell(row=row, column=serial_col)
        val = cell.value
        if _is_formula_cell(cell):
            수식발견 = True
            빈행 = 0
            continue
        if val in (None, ""):
            빈행 += 1
            if 빈행 >= _EMPTY_RUN_STOP:
                break
            continue
        빈행 = 0
        try:
            last = max(last, int(val))
        except (TypeError, ValueError):
            pass
    if 수식발견 and last == 0:
        return -1     # 연번 열이 수식 → 직접 채번하지 않음
    return last


def _copy_style(ws, src: int, dst: int):
    for col in range(1, ws.max_column + 1):
        s = ws.cell(row=src, column=col)
        d = ws.cell(row=dst, column=col)
        if s.has_style:
            d.font      = copy(s.font)
            d.fill      = copy(s.fill)
            d.border    = copy(s.border)
            d.alignment = copy(s.alignment)
            d.number_format = s.number_format
    if ws.row_dimensions[src].height:
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height


# ─── 헤더 자동 분석 (새 사무소 매핑 생성 보조) ──────────────────────────────

# 헤더 키워드 → 필드명 매핑
HEADER_KEYWORDS = {
    # 기본정보
    "연번":            ["연번", "번호", "순번", "no", "번"],
    "서류수령일":      ["수령일", "서류수령", "접수일", "수령"],
    "동":              ["동", "동번호", "동호수", "아파트동"],
    "호":              ["호", "호수", "호번호", "아파트호"],
    "성명":            ["성명", "이름", "수분양자", "매수인", "소유자", "성함",
                        "매수인성명", "수분양자성명"],
    "전화번호":        ["연락처", "전화번호", "전화", "휴대폰", "휴대전화",
                        "핸드폰", "전화(휴대)"],
    "주민등록번호":    ["주민등록번호", "주민번호", "생년월일", "주민"],
    "주소":            ["주소", "현주소", "주소지", "거주지", "거주주소",
                        "현재주소"],
    "전용면적":        ["전용면적", "전용", "건물면적", "면적", "㎡"],
    "대지지분":        ["대지지분", "대지권", "대지면적", "대지"],
    "초본":            ["초본", "주민등록초본", "초본발행"],
    "인감":            ["인감", "인감증명", "인감발행"],
    "등본":            ["등본", "주민등록등본", "등본발행"],
    "미비서류":        ["미비서류", "미비", "미비사항", "미비내용"],
    # 분양정보
    "분양계약일":      ["계약일", "분양계약일", "계약날짜", "분양계약"],
    "분양대금":        ["분양가격", "분양대금", "분양가", "총분양가",
                        "분양금액", "매매대금"],
    "부가세":          ["부가세", "vat", "부가가치세"],
    "발코니금액":      ["발코니", "발코니금액", "발코니확장"],
    "옵션금액":        ["옵션", "유상옵션", "옵션금액", "선택품목"],
    "거래가액":        ["거래가액", "실거래가", "거래금액", "매매가"],
    # 승계
    "승계여부":        ["승계여부", "권리의무승계", "승계", "권리승계"],
    "승계일":          ["승계일", "승계날짜", "권리이전일"],
    "거래신고필증번호": ["신고필증", "거래신고", "필증번호", "신고번호"],
    # 대출/채권
    "대출은행":        ["대출은행", "은행", "금융기관"],
    "대출지점":        ["지점", "대출지점", "은행지점"],
    "채권최고액":      ["채권최고액", "채권액", "최고채권"],
    "근저당설정계약일": ["설정계약일", "근저당설정", "설정일"],
    # 취득세/비용
    "취득세":          ["취득세", "취득세액"],
    "교육세":          ["교육세"],
    "농특세":          ["농특세", "농어촌특별세"],
    "취득세합계":      ["취득세합계", "세금합계", "취득관련세금"],
    "등기비용총합계":  ["등기비용합계", "등기비용", "총등기비용", "비용합계"],
}


def analyze_template(xlsx_path: str, 사무소명: str,
                     header_row: int = None) -> dict:
    """
    템플릿 엑셀의 헤더를 분석해서 매핑 자동 생성.
    header_row=None 이면 1~5행 중 매칭 가장 많은 행 자동 선택.
    """
    wb = openpyxl.load_workbook(xlsx_path)

    # 시트 탐색: 기본명단 포함 > 첫 번째
    ws = next((wb[s] for s in wb.sheetnames
               if any(k in s for k in ["기본명단","명단","APT","아파트","오피"])),
              wb.active)
    sheet_name = ws.title

    # 헤더행 자동 탐지
    if header_row is None:
        best_row, best_count = 1, 0
        for r in range(1, 6):
            count = 0
            for col in range(1, min(ws.max_column+1, 220)):
                v = str(ws.cell(row=r, column=col).value or "").replace("\n","").replace(" ","")
                for kws in HEADER_KEYWORDS.values():
                    for kw in kws:
                        if kw.replace(" ","") in v:
                            count += 1
                            break
            if count > best_count:
                best_count, best_row = count, r
        header_row = best_row

    columns = {}
    for col in range(1, ws.max_column + 1):
        cell_val = str(ws.cell(row=header_row, column=col).value or "").strip()
        cell_val_clean = cell_val.replace("\n", "").replace(" ", "")
        for field, keywords in HEADER_KEYWORDS.items():
            for kw in keywords:
                if kw.replace(" ", "") in cell_val_clean:
                    if field not in columns:
                        columns[field] = col
                    break

    mapping = {
        "_info": {
            "사무소명":   사무소명,
            "템플릿":     Path(xlsx_path).name,
            "시트명":     sheet_name,
            "헤더행":     header_row,
            "데이터시작행": header_row + 1,
        },
        "_columns":        columns,
        "_amount_columns": [],
        "_key_column":     columns.get("성명", 1),
    }
    return mapping
