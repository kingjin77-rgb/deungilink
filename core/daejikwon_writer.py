"""
대지권 주소명단 자동 입력 모듈
registry_auto self._dj_results → 주소명단 시트 자동 기입

컬럼 매핑 (교차검증 완료 2026-05-13):
  A(1)  연번              B(2)  비고
  C(3)  아파트명칭        D(4)  개별/공동
  E(5)  동               F(6)  호수
  G(7)  성명             H(8)  주소
  I(9)  건물등기 접수일자  J(10) 건물면적
  K(11) 대지권면적        L(12) * (황색 보조열 — 입력 금지)
  M(13) 동호수 (수식)
"""
import copy
import re
from collections import Counter
from datetime import date, datetime
from openpyxl.utils import get_column_letter
from openpyxl.styles import numbers as xl_num


# ── 헬퍼 ───────────────────────────────────────────────────────────────────

def _s(v):
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v
    return str(v).strip()


def clean_address(addr):
    if not addr:
        return addr
    addr = re.sub(r'^[\s\-]*[\*#]+[\s\*#\-]*', '', addr)  # 앞 마스킹 제거
    addr = addr.strip(' |')                                 # 뒤 | 제거
    return addr.strip()


def _float(v):
    try:
        return float(_s(v)) if _s(v) else None
    except (ValueError, TypeError):
        return None


def _to_date(v):
    """문자열 "YYYY-MM-DD" 또는 date 객체 → date 객체. 실패 시 원본 반환."""
    if isinstance(v, (date, datetime)):
        return v if isinstance(v, date) else v.date()
    s = str(v).strip()
    try:
        return date.fromisoformat(s)     # "2026-03-15"
    except (ValueError, AttributeError):
        pass
    import re
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    return s   # 파싱 불가 시 원래 문자열 유지


def _fmt_dong(v):
    v = _s(v)
    return (v + "동") if v and not v.endswith("동") else v


def _fmt_ho(v):
    v = _s(v)
    return (v + "호") if v and not v.endswith("호") else v


def _detect_bigo(rec, standard_dj):
    """비고 자동 판단 (우선순위: 외국인 > 미분양 > 지분상이)"""
    if _s(rec.get("국적", "")) and any(k in _s(rec.get("국적", ""))
                                        for k in ["외국", "foreign", "Foreign"]):
        return "외국인"
    name = _s(rec.get("성명", ""))
    if "미분양" in _s(rec.get("분양여부", "")) or not name:
        return "미분양"
    if standard_dj is not None:
        dj = _float(rec.get("대지지분") or rec.get("대지권면적"))
        if dj is not None and abs(dj - standard_dj) > 0.001:
            return "지분상이"
    return ""


def _build_names(rec):
    n1 = _s(rec.get("성명", ""))
    n2 = _s(rec.get("공동명의자", ""))
    return f"{n1}, {n2}" if n2 else n1


def _copy_style(src, dst):
    dst.font         = copy.copy(src.font)
    dst.border       = copy.copy(src.border)
    dst.fill         = copy.copy(src.fill)
    dst.number_format = src.number_format
    dst.alignment    = copy.copy(src.alignment)


# ── 메인 함수 ───────────────────────────────────────────────────────────────

def write_address_sheet(ws, records: list, apt_name: str = "",
                        building_date=None):
    """
    주소명단 워크시트에 데이터 입력.
    기존 서식(테두리, 폰트, 정렬) 완전 보존.

    Parameters
    ----------
    ws            : openpyxl Worksheet ('주소명단')
    records       : list[dict] — self._dj_results
    apt_name      : 아파트 명칭 (예: '부영3단지')
    building_date : 건물등기 접수일자 공통값 (개별 값 없을 때 사용)
    """
    DATA_START = 2   # 1행 헤더, 2행부터 데이터

    # 표준 대지권 자동 산출 (최빈값 → 지분상이 판단 기준)
    dj_vals = []
    for r in records:
        v = _float(r.get("대지지분") or r.get("대지권면적"))
        if v:
            dj_vals.append(round(v, 4))
    standard_dj = Counter(dj_vals).most_common(1)[0][0] if dj_vals else None

    ref_row = DATA_START   # 서식 복사 기준행

    # 동/호/성명 모두 공란인 빈 행 제거
    records = [r for r in records
               if _s(r.get("동")) or _s(r.get("호")) or _s(r.get("성명"))]

    for i, rec in enumerate(records):
        row = DATA_START + i

        # 서식 복사 (2행 기준 → 모든 데이터행 동일 적용)
        if row != ref_row:
            for col in range(1, ws.max_column + 1):
                _copy_style(ws.cell(ref_row, col), ws.cell(row, col))
            if ref_row in ws.row_dimensions:
                ws.row_dimensions[row].height = ws.row_dimensions[ref_row].height

        # ── 값 계산 ────────────────────────────────────────────────────────
        apt      = apt_name or _s(rec.get("아파트명칭", ""))
        dong     = _fmt_dong(rec.get("동", ""))
        ho       = _fmt_ho(rec.get("호", "") or rec.get("호수", ""))
        own      = "공동" if _s(rec.get("공동명의자", "")) else "개별"
        bigo     = _detect_bigo(rec, standard_dj)
        names    = _build_names(rec)
        addr     = clean_address(_s(rec.get("주소", "")))
        건물면적  = _float(rec.get("전용면적") or rec.get("건물면적"))

        # 건물등기 접수일자: 개별 값 우선, 없으면 공통값
        접수일_raw = rec.get("건물등기접수일자") or building_date
        접수일     = _to_date(접수일_raw) if 접수일_raw else None

        # ── 셀 입력 (교차검증 완료) ────────────────────────────────────────
        # A: 연번
        c = ws.cell(row, 1)
        c.value = i + 1

        # B: 비고
        c = ws.cell(row, 2)
        c.value = bigo

        # C: 아파트명칭
        c = ws.cell(row, 3)
        c.value = apt

        # D: 개별/공동
        c = ws.cell(row, 4)
        c.value = own

        # E: 동
        c = ws.cell(row, 5)
        c.value = dong

        # F: 호수
        c = ws.cell(row, 6)
        c.value = ho

        # G: 성명
        c = ws.cell(row, 7)
        c.value = names

        # H: 주소
        c = ws.cell(row, 8)
        c.value = addr

        # I: 건물등기 접수일자 (date 객체로 입력 → Excel 날짜 서식)
        c = ws.cell(row, 9)
        c.value = 접수일
        if isinstance(접수일, date):
            c.number_format = "YYYY-MM-DD"

        # J: 건물면적 (숫자)
        c = ws.cell(row, 10)
        c.value = 건물면적

        # K: 대지권면적 — 수기 입력 (공란 유지)

        # L: * 황색 보조열 — 절대 입력하지 않음

        # M: 동호수 수식 — 동이 공란이면 "-호수" 방지
        E = get_column_letter(5)   # E
        F = get_column_letter(6)   # F
        ws.cell(row, 13).value = f'=IF({E}{row}="","",{E}{row}&"-"&{F}{row})'

    # 남은 템플릿 행 초기화 (데이터 없는 행)
    E = get_column_letter(5)
    F = get_column_letter(6)
    for row in range(DATA_START + len(records), ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if   col == 1:  ws.cell(row, col).value = row - DATA_START + 1
            elif col == 4:  ws.cell(row, col).value = "개별"
            elif col == 13: ws.cell(row, col).value = f'=IF({E}{row}="","",{E}{row}&"-"&{F}{row})'
            elif col != 12: ws.cell(row, col).value = None
