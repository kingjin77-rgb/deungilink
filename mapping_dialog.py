"""
사무소 매핑 추가/편집 다이얼로그
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import json
from pathlib import Path
from core.mapping_manager import (
    load_mapping, save_mapping, delete_mapping,
    analyze_template, MAPPINGS_DIR
)

C = {
    "bg": "#0F1117", "panel": "#181C27", "card": "#1E2436",
    "border": "#2A3050", "accent": "#3B82F6", "green": "#22C55E",
    "red": "#EF4444", "yellow": "#F59E0B",
    "text": "#E2E8F0", "sub": "#64748B", "white": "#FFFFFF",
}
FONT  = ("맑은 고딕", 9)
FONTB = ("맑은 고딕", 9, "bold")
FONTH = ("맑은 고딕", 11, "bold")
FONTS = ("맑은 고딕", 8)

# 입력 가능한 전체 필드 목록
ALL_FIELDS = [
    ("기본정보", [
        ("연번",          "A열 순번"),
        ("서류수령일",    "D열 서류 수령일"),
        ("동",            "아파트 동"),
        ("호",            "아파트 호"),
        ("성명",          "매수인 성명"),
        ("전화번호",      "연락처"),
        ("주민등록번호",  "주민등록번호"),
        ("주소",          "현주소"),
        ("전용면적",      "전용면적(㎡)"),
        ("대지지분",      "대지지분(㎡)"),
        ("초본",          "초본 발행일"),
        ("인감",          "인감 발행일"),
        ("미비서류",      "미비서류 내용"),
    ]),
    ("분양정보", [
        ("분양계약일",    "분양계약일"),
        ("분양대금",      "총 분양대금"),
        ("부가세",        "부가가치세"),
        ("발코니금액",    "발코니 금액"),
        ("옵션금액",      "유상옵션 금액"),
    ]),
    ("권리의무승계", [
        ("승계여부",          "승계여부"),
        ("승계일",            "승계일자"),
        ("거래신고필증번호",  "거래신고필증번호"),
        ("거래가액",          "거래가액"),
    ]),
    ("취득세", [
        ("취득세과표",  "BL 취득세과표"),
        ("취득세",      "BM 취득세"),
        ("교육세",      "BN 교육세"),
        ("농특세",      "BO 농특세"),
        ("취득세합계",  "BP 취득세합계"),
    ]),
    ("등기비용", [
        ("등기비용총합계",     "CE 등기비용합계"),
        ("등기비용합계_AL",    "AL 등기비용합계(이중)"),
        ("채권매입금액_이전",  "BQ 이전채권매입"),
        ("채권할인금액_이전",  "BR 이전채권할인"),
        ("인지대_이전",        "BS 인지대"),
        ("증지대_이전",        "BT 증지대"),
        ("보수료",             "BX 이전보수료"),
        ("부가세_이전",        "CA 부가세"),
        ("신탁말소비용",       "CC 신탁말소"),
        ("설정비용합계",       "CB 설정비용청구"),
        ("등록세_설정",        "CS 설정등록세"),
        ("교육세_설정",        "CT 설정교육세"),
        ("채권매입금액_설정",  "CU 설정채권매입"),
        ("채권할인금액_설정",  "CV 설정채권할인"),
        ("보수료_설정",        "DC 설정보수료"),
        ("부가세_설정",        "DF 설정부가세"),
        ("설정비용1순위",      "DI 설정비용합계"),
    ]),
    ("대출정보", [
        ("대출은행",          "CF 대출은행"),
        ("대출지점",          "CG 지점"),
        ("채권최고액",        "CO 채권최고액"),
        ("근저당설정계약일",  "CR 설정계약일"),
    ]),
]


class MappingDialog(tk.Toplevel):
    def __init__(self, parent, 사무소명: str = None):
        super().__init__(parent)
        self._edit_mode = 사무소명 is not None
        self._사무소명  = 사무소명
        self._col_vars  = {}  # field → StringVar (열 번호)

        self.title("⚙  매핑 편집" if self._edit_mode else "➕  새 사무소 추가")
        self.geometry("620x700")
        self.resizable(False, True)
        self.configure(bg=C["bg"])
        self.grab_set()

        self._build()
        if self._edit_mode:
            self._load_existing()

    def _build(self):
        # 헤더
        hdr = tk.Frame(self, bg=C["panel"], height=48)
        hdr.pack(fill="x")
        hdr.pack_propagate(False)
        title = "매핑 편집" if self._edit_mode else "새 사무소 추가"
        tk.Label(hdr, text=f"🏢  {title}", font=FONTH,
                 fg=C["accent"], bg=C["panel"]).pack(side="left", padx=16, pady=10)

        # 사무소명 입력
        nf = tk.Frame(self, bg=C["bg"])
        nf.pack(fill="x", padx=20, pady=(12, 4))
        tk.Label(nf, text="사무소명", font=FONTB, fg=C["text"], bg=C["bg"]).pack(anchor="w")
        self._name_var = tk.StringVar(value=self._사무소명 or "")
        name_ent = tk.Entry(nf, textvariable=self._name_var, font=FONT,
                            bg=C["card"], fg=C["text"], insertbackground=C["text"],
                            relief="flat", bd=6)
        name_ent.pack(fill="x", ipady=5)
        if self._edit_mode:
            name_ent.config(state="readonly")

        # 자동분석 버튼
        af = tk.Frame(self, bg=C["bg"])
        af.pack(fill="x", padx=20, pady=(4, 8))
        tk.Button(af, text="📂  엑셀 템플릿으로 자동 분석",
                  command=self._auto_analyze,
                  font=FONTB, bg=C["accent"], fg=C["white"],
                  relief="flat", padx=12, pady=6, cursor="hand2").pack(side="left")
        tk.Label(af, text="← 템플릿 열면 열번호 자동 입력",
                 font=FONTS, fg=C["sub"], bg=C["bg"]).pack(side="left", padx=8)

        # 구분선
        tk.Frame(self, bg=C["border"], height=1).pack(fill="x", padx=20)

        # 스크롤 영역 (필드 입력)
        canvas = tk.Canvas(self, bg=C["bg"], highlightthickness=0)
        sb = tk.Scrollbar(self, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        frame = tk.Frame(canvas, bg=C["bg"])
        win = canvas.create_window((0, 0), window=frame, anchor="nw")

        def _resize(e):
            canvas.configure(scrollregion=canvas.bbox("all"))
            canvas.itemconfig(win, width=canvas.winfo_width())
        frame.bind("<Configure>", _resize)
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(win, width=e.width))
        canvas.bind_all("<MouseWheel>",
                        lambda e: canvas.yview_scroll(-1*(e.delta//120), "units"))

        # 섹션별 필드
        for section_name, fields in ALL_FIELDS:
            sf = tk.Frame(frame, bg=C["bg"])
            sf.pack(fill="x", padx=16, pady=(12, 2))
            tk.Label(sf, text=section_name, font=FONTB,
                     fg=C["accent"], bg=C["bg"]).pack(side="left")
            tk.Frame(sf, bg=C["border"], height=1).pack(
                side="left", fill="x", expand=True, padx=(6,0), pady=4)

            for field, desc in fields:
                row = tk.Frame(frame, bg=C["card"])
                row.pack(fill="x", padx=16, pady=1)
                tk.Label(row, text=desc, font=FONT, fg=C["text"],
                         bg=C["card"], width=22, anchor="w").pack(side="left", padx=8, pady=5)
                var = tk.StringVar()
                tk.Entry(row, textvariable=var, font=FONT,
                         bg=C["bg"], fg=C["yellow"], insertbackground=C["text"],
                         relief="flat", bd=4, width=6,
                         justify="center").pack(side="left", padx=4, ipady=3)
                tk.Label(row, text="열번호", font=FONTS,
                         fg=C["sub"], bg=C["card"]).pack(side="left")
                self._col_vars[field] = var

        # 하단 버튼
        bf = tk.Frame(frame, bg=C["bg"])
        bf.pack(fill="x", padx=16, pady=16)
        tk.Button(bf, text="✅  저장", command=self._save,
                  font=FONTB, bg=C["green"], fg=C["white"],
                  relief="flat", padx=20, pady=8, cursor="hand2").pack(side="left")
        if self._edit_mode:
            tk.Button(bf, text="🗑  삭제", command=self._delete,
                      font=FONT, bg=C["red"], fg=C["white"],
                      relief="flat", padx=14, pady=8, cursor="hand2").pack(side="left", padx=8)
        tk.Button(bf, text="닫기", command=self.destroy,
                  font=FONT, bg=C["panel"], fg=C["sub"],
                  relief="flat", padx=14, pady=8, cursor="hand2").pack(side="right")

    def _load_existing(self):
        """기존 매핑 로드"""
        try:
            mapping = load_mapping(self._사무소명)
            col_map = mapping.get("_columns", {})
            for field, var in self._col_vars.items():
                if field in col_map:
                    var.set(str(col_map[field]))
        except Exception as e:
            messagebox.showerror("오류", f"매핑 로드 실패: {e}", parent=self)

    def _auto_analyze(self):
        """템플릿 엑셀 자동 분석 — 결과 미달 시 실제 헤더 목록 표시"""
        path = filedialog.askopenfilename(
            title="기본명단 템플릿 선택",
            filetypes=[("Excel", "*.xlsx")], parent=self)
        if not path:
            return
        try:
            사무소명 = self._name_var.get().strip() or "임시"
            mapping = analyze_template(path, 사무소명)
            col_map = mapping.get("_columns", {})
            found = 0
            for field, var in self._col_vars.items():
                if field in col_map:
                    var.set(str(col_map[field]))
                    found += 1

            if found == 0:
                # 실제 헤더 목록 읽어서 표시
                import openpyxl
                wb = openpyxl.load_workbook(path, data_only=True)
                ws = next((wb[s] for s in wb.sheetnames if "기본명단" in s), wb.active)
                headers = []
                for col in range(1, ws.max_column + 1):
                    v = ws.cell(row=1, column=col).value
                    if v:
                        headers.append(f"{col}열: {str(v).strip()[:15]}")

                header_txt = "\n".join(headers[:30]) if headers else "헤더 없음"
                messagebox.showwarning(
                    "자동 인식 0개",
                    f"헤더명이 매핑 키워드와 다릅니다.\n\n"
                    f"[엑셀 1행 헤더 목록]\n{header_txt}\n\n"
                    f"위 열번호를 참고하여 직접 입력하세요.",
                    parent=self)
            else:
                messagebox.showinfo("자동 분석 완료",
                                    f"총 {found}개 열 자동 인식 완료!\n"
                                    f"나머지는 직접 입력해주세요.", parent=self)
        except Exception as e:
            messagebox.showerror("오류", str(e), parent=self)

    def _save(self):
        사무소명 = self._name_var.get().strip()
        if not 사무소명:
            messagebox.showwarning("경고", "사무소명을 입력해주세요.", parent=self)
            return

        col_map = {}
        amount_cols = []
        AMOUNT_FIELDS = {"분양대금","부가세","발코니금액","옵션금액","거래가액",
                         "취득세과표","취득세","교육세","농특세","취득세합계",
                         "채권최고액","등기비용총합계","등기비용합계_AL",
                         "채권매입금액_이전","채권할인금액_이전","인지대_이전",
                         "보수료","부가세_이전","신탁말소비용","설정비용합계",
                         "등록세_설정","교육세_설정","채권매입금액_설정",
                         "채권할인금액_설정","보수료_설정","부가세_설정",
                         "설정비용1순위","거래가액"}

        for field, var in self._col_vars.items():
            val = var.get().strip()
            if val:
                try:
                    col_num = int(val)
                    col_map[field] = col_num
                    if field in AMOUNT_FIELDS:
                        amount_cols.append(col_num)
                except ValueError:
                    pass

        if not col_map:
            messagebox.showwarning("경고", "최소 1개 이상 열번호를 입력해주세요.", parent=self)
            return

        # 기존 매핑 기반 info 유지 or 새로 생성
        if self._edit_mode:
            try:
                existing = load_mapping(사무소명)
                info = existing.get("_info", {})
            except:
                info = {}
        else:
            info = {}

        info["사무소명"] = 사무소명

        mapping = {
            "_info":            info,
            "_columns":         col_map,
            "_amount_columns":  sorted(set(amount_cols)),
            "_key_column":      col_map.get("성명", 1),
        }
        save_mapping(사무소명, mapping)
        messagebox.showinfo("저장 완료", f"[{사무소명}] 매핑 저장 완료!", parent=self)
        self.destroy()

    def _delete(self):
        if messagebox.askyesno("삭제 확인",
                               f"[{self._사무소명}] 매핑을 삭제할까요?", parent=self):
            delete_mapping(self._사무소명)
            self.destroy()
